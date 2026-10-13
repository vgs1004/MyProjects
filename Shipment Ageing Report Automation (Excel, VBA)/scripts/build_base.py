"""
Builds the base workbook (all 6 sheets, Settings values, Dashboard KPI
cards, Control/Run Log layout) with openpyxl. VBA modules, the
Control-sheet buttons, and the Dashboard charts are added afterwards via
Excel COM (inject_and_run.py), since openpyxl can't touch the VBA
project or create chart/button objects that VBA can still manage.
"""
import datetime
import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---- palette: clean neutral ink + a vibrant flat-colour KPI-card set,
# modern dashboard look (not the muted earth tones used elsewhere in
# this portfolio -- matched to the examples the project asked for) ----
INK = "1F2430"         # near-black navy-charcoal, headers/body text
LABEL_GRAY = "7A8699"
CARD_BORDER = "E3E6EC"

CARD_BLUE = "4361EE"    # Exception Lines
CARD_PINK = "F72585"    # % Late
CARD_PURPLE = "7209B7"  # Average Delay
CARD_ORANGE = "F77F00"  # Share 3+ Days Late
CARD_TEAL = "2A9D8F"    # Top Region

MET_FILL, MET_FONT = "DCE7D2", "4B6142"     # kept for Exceptions RAG formatting
OPEN_FILL, OPEN_FONT = "F5E7C9", "8C6A22"
MID_FILL, MID_FONT = "F0CDAA", "994C00"
BREACH_FILL, BREACH_FONT = "F0DCD6", "8C4B3A"

REGIONS = [
    "Canada", "Caribbean", "Central Africa", "Central America", "Central Asia",
    "East Africa", "East of USA", "Eastern Asia", "Eastern Europe", "North Africa",
    "Northern Europe", "Oceania", "South America", "South Asia", "South of USA",
    "Southeast Asia", "Southern Africa", "Southern Europe", "US Center",
    "West Africa", "West Asia", "West of USA", "Western Europe",
]


def slug(region):
    return region.lower().replace(" ", "")


wb = openpyxl.Workbook()
header_fill = PatternFill("solid", fgColor=INK)
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color=CARD_BORDER)
box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

# ============================== Control ==============================
ws = wb.active
ws.title = "Control"
ws.sheet_view.showGridLines = False
ws["D2"] = "Shipment Ageing Report Automation"
ws["D2"].font = Font(size=16, bold=True, color=INK)
ws["D4"] = "Buttons are added to the right of this block once the macros are loaded."
ws["D4"].font = Font(size=9, italic=True, color=LABEL_GRAY)

status_card_fill = PatternFill("solid", fgColor="F4F6FA")
status_rows = [(2, "Last Run"), (3, "Exception Lines"), (4, "Runtime"), (5, "Status"), (6, "Error Detail")]
for row, lab in status_rows:
    ws.cell(row=row, column=1, value=lab).font = Font(bold=True, color=INK)
    ws.cell(row=row, column=2).font = Font(color=CARD_BLUE, bold=True)
    for c in (1, 2):
        ws.cell(row=row, column=c).fill = status_card_fill
        ws.cell(row=row, column=c).border = box_border
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 3
ws.column_dimensions["D"].width = 36

# ============================== Settings ==============================
ws = wb.create_sheet("Settings")
ws.sheet_view.showGridLines = False
ws["A1"] = "SETTINGS"
ws["A1"].font = Font(size=14, bold=True, color=INK)

settings_rows = [
    (3, "Report Date", datetime.date(2018, 2, 6)),
    (4, "Lookback Days", 30),
    (6, "Excluded Order Statuses (comma, no spaces)", "CANCELED,SUSPECTED_FRAUD,COMPLETE,CLOSED"),
    (7, "Excluded Delivery Status", "Shipping canceled"),
    (9, "On-schedule max delay (days)", 0),
    (10, "Day-1 bucket max delay", 1),
    (11, "Day-2 bucket max delay", 2),
    (13, "Email Subject", "Daily Shipment Ageing Report"),
    (14, "Email Body", "Please find the pending last-mile shipments for your region. "
                        "We request you to proceed with these on priority."),
    (16, "Draft Emails on Run (Yes/No)", "No"),
]
for row, label, value in settings_rows:
    ws.cell(row=row, column=1, value=label).font = Font(color=INK)
    c = ws.cell(row=row, column=2, value=value)
    c.font = Font(bold=True, color=CARD_BLUE)
    if isinstance(value, datetime.date):
        c.number_format = "yyyy-mm-dd"

ws["A17"] = "Region"
ws["B17"] = "Email"
for cell in ("A17", "B17"):
    ws[cell].font = header_font
    ws[cell].fill = header_fill
for i, region in enumerate(REGIONS, start=18):
    ws.cell(row=i, column=1, value=region)
    ws.cell(row=i, column=2, value=f"{slug(region)}@example.com")

ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 40

# ============================== RawData ==============================
ws = wb.create_sheet("RawData")
ws["A1"] = "Populated by Run Daily Report / Re-import CSV. Safe to ignore between runs."
ws["A1"].font = Font(size=9, italic=True, color=LABEL_GRAY)

# ============================== Exceptions ==============================
ws = wb.create_sheet("Exceptions")
ws["A1"] = "Populated by Run Daily Report."
ws["A1"].font = Font(size=9, italic=True, color=LABEL_GRAY)

# ============================== Dashboard ==============================
# Replaces the old separate Pivots + Summary sheets with one real
# dashboard: 5 colour KPI cards (openpyxl, live formulas) plus 4 native
# charts and a top-cities table (VBA, see modBuildDashboard.bas) --
# charts instead of raw pivot grids, one sheet instead of two.
ws = wb.create_sheet("Dashboard")
ws.sheet_view.showGridLines = False
ws["B2"] = "Shipment Ageing Dashboard"
ws["B2"].font = Font(size=18, bold=True, color=INK)

kpi_defs = [
    ("Exception Lines", '=COUNTIFS(Exceptions!A2:A1048576,"<>")', None, CARD_BLUE),
    ("% Late", '=IFERROR(COUNTIFS(Exceptions!R2:R1048576,"<>On schedule",Exceptions!A2:A1048576,"<>")/$B$5,0)', "0.0%", CARD_PINK),
    ("Avg Delay (days)", '=IFERROR(AVERAGEIFS(Exceptions!Q2:Q1048576,Exceptions!A2:A1048576,"<>"),0)', "0.00", CARD_PURPLE),
    ("Share 3+ Days Late", '=IFERROR(COUNTIFS(Exceptions!R2:R1048576,"3+ days late",Exceptions!A2:A1048576,"<>")/$B$5,0)', "0.0%", CARD_ORANGE),
    ("Top Region", "", None, CARD_TEAL),  # written as a plain value by modBuildDashboard.bas
]

start_col = 2
for i, (label, formula, fmt, color) in enumerate(kpi_defs):
    col = start_col + i * 3  # 2 cols per card + 1 spacer col between cards
    col_l = get_column_letter(col)
    col_l2 = get_column_letter(col + 1)
    fill = PatternFill("solid", fgColor=color)
    center = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"{col_l}4:{col_l2}4")
    ws.merge_cells(f"{col_l}5:{col_l2}5")
    lbl = ws[f"{col_l}4"]
    lbl.value = label.upper()
    lbl.font = Font(size=9, color="FFFFFF", bold=True)
    lbl.alignment = center
    val = ws[f"{col_l}5"]
    if formula:
        val.value = formula
        if fmt:
            val.number_format = fmt
    val.font = Font(size=20, bold=True, color="FFFFFF")
    val.alignment = center
    for r in (4, 5, 6):
        ws.cell(row=r, column=col).fill = fill
        ws.cell(row=r, column=col + 1).fill = fill
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 32
    ws.row_dimensions[6].height = 6
    ws.column_dimensions[col_l].width = 13
    ws.column_dimensions[get_column_letter(col + 1)].width = 10

ws["B8"] = "Charts below are built from the Exceptions sheet each run. KPI values are live formulas."
ws["B8"].font = Font(size=9, italic=True, color=LABEL_GRAY)

# ============================== Run Log ==============================
ws = wb.create_sheet("Run Log")
headers = ["Date Run", "Report Date", "Rows In", "Rows Out", "Seconds", "Drafts Created"]
ws.append(headers)
for col in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col)
    c.fill = header_fill
    c.font = header_font
widths = [20, 14, 10, 10, 10, 14]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

out_path = os.path.join(HERE, "Shipment_Ageing_Report_base.xlsx")
wb.save(out_path)
print("Saved", out_path)
